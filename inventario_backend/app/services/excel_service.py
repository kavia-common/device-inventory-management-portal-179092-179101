from __future__ import annotations

"""
Excel import/export service for InventoryItem using openpyxl.

Provides utilities to:
- Parse a .xlsx file with a header row and produce row dicts + validation errors
- Generate a .xlsx workbook from a list of item dicts for download
"""

from io import BytesIO
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


# PUBLIC_INTERFACE
def parse_items_from_excel(file_stream) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Parse inventory items from an uploaded Excel (.xlsx) file.

    The file must have a header row with recognizable columns. Supported column
    names (case-insensitive, spaces/underscores tolerated):

    - device_type (required)
    - brand
    - model
    - serial_number
    - owner
    - location
    - status
    - notes
    - asset_tag (optional, even if not persisted in model—schema may validate/ignore)

    Args:
        file_stream: A file-like object pointing to an Excel .xlsx file.

    Returns:
        tuple:
            - items: list of dicts suitable for creating/updating InventoryItem
            - errors: list of row error dicts: {"row": <row_index>, "message": <str>}
    """
    # Load workbook
    wb = load_workbook(filename=file_stream, read_only=True, data_only=True)
    ws = wb.active  # type: Worksheet

    # Normalize header map
    header_map = _extract_header_map(ws)
    items: List[Dict[str, Any]] = []
    errors: List[Dict[str, Any]] = []

    required_cols = ["device_type"]
    # Iterate rows, starting after header
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        row_dict = _row_to_item_dict(header_map, row)
        # Skip completely empty rows
        if not any(v is not None and str(v).strip() != "" for v in row_dict.values()):
            continue

        # Validate required columns presence
        missing = [c for c in required_cols if not _has_value(row_dict.get(c))]
        if missing:
            errors.append({"row": idx, "message": f"Missing required fields: {', '.join(missing)}"})
            continue

        # Normalize strings and strip whitespace
        for k, v in list(row_dict.items()):
            if isinstance(v, str):
                row_dict[k] = v.strip()

        items.append(row_dict)

    return items, errors


def _extract_header_map(ws: Worksheet) -> Dict[int, str]:
    """Build a header map: column_index -> normalized_field_name."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    mapping: Dict[int, str] = {}
    for i, name in enumerate(header_row):
        if not name:
            continue
        normalized = _normalize_header(str(name))
        if normalized:
            # map column index to normalized field name
            mapping[i] = normalized
    return mapping


def _normalize_header(name: str) -> Optional[str]:
    """Normalize column header names to our canonical field keys."""
    key = name.strip().lower().replace(" ", "_")
    # accepted aliases map
    aliases = {
        "device_type": ["device_type", "type", "tipo", "tipo_dispositivo"],
        "brand": ["brand", "marca"],
        "model": ["model", "modelo"],
        "serial_number": ["serial_number", "serial", "no_serie", "n_serie", "número_de_serie"],
        "owner": ["owner", "assigned_to", "asignado_a", "propietario"],
        "location": ["location", "ubicacion", "ubicación"],
        "status": ["status", "estado"],
        "notes": ["notes", "notas", "comentarios"],
        "asset_tag": ["asset_tag", "activo", "codigo_activo", "código_activo"],
    }
    for canonical, names in aliases.items():
        if key in names:
            return canonical
    # If the key already matches one of canonical names, accept as-is
    if key in aliases:
        return key
    return None


def _row_to_item_dict(header_map: Dict[int, str], row: Tuple[Any, ...]) -> Dict[str, Any]:
    """Create a dict from a row tuple using the header_map."""
    item: Dict[str, Any] = {
        "device_type": None,
        "brand": None,
        "model": None,
        "serial_number": None,
        "owner": None,
        "location": None,
        "status": None,
        "notes": None,
        "asset_tag": None,
    }
    for col_idx, field_name in header_map.items():
        if col_idx < len(row):
            item[field_name] = row[col_idx]
    return item


def _has_value(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, str) and value.strip() == "":
        return False
    return True


# PUBLIC_INTERFACE
def generate_items_excel(items: Iterable[Dict[str, Any]]) -> bytes:
    """Generate an Excel (.xlsx) file bytes from an iterable of item dicts.

    Columns output in this order:
    device_type, brand, model, serial_number, owner, location, status, notes, asset_tag, created_at, updated_at, id

    Args:
        items: Iterable of item dictionaries (e.g., from InventoryItem.to_dict()).

    Returns:
        bytes: The binary content of the generated .xlsx file.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"

    headers = [
        "device_type",
        "brand",
        "model",
        "serial_number",
        "owner",
        "location",
        "status",
        "notes",
        "asset_tag",
        "created_at",
        "updated_at",
        "id",
    ]
    ws.append(headers)

    for item in items:
        row = [
            item.get("device_type"),
            item.get("brand"),
            item.get("model"),
            item.get("serial_number"),
            item.get("owner"),
            item.get("location"),
            item.get("status"),
            item.get("notes"),
            item.get("asset_tag"),
            item.get("created_at"),
            item.get("updated_at"),
            item.get("id"),
        ]
        ws.append(row)

    # Autosize columns (roughly)
    for col in ws.columns:
        max_len = 10
        col_letter = col[0].column_letter  # type: ignore[attr-defined]
        for cell in col:
            val = cell.value
            if val is None:
                continue
            length = len(str(val))
            if length > max_len:
                max_len = length
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)

    stream = BytesIO()
    wb.save(stream)
    return stream.getvalue()
